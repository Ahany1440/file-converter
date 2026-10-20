"""Runs every conversion in the map on a sample file and checks the result opens."""
import json
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import converters  # noqa: E402
from converters.tools import ConversionError, find, run, status  # noqa: E402
from tests.make_samples import make  # noqa: E402

TOOLS = status()


@pytest.fixture(scope="session")
def samples():
    return make()


def check(path):
    """Open the output with the library for its type. Fails the test if it can't."""
    assert path.exists(), f"{path.name} was not created"
    assert path.stat().st_size > 0, f"{path.name} is empty"
    ext = path.suffix.lower().lstrip(".")
    if ext == "zip":
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
            assert names
            z.extractall(path.parent / "_unzipped")
        for n in names:
            check(path.parent / "_unzipped" / n)
    elif ext == "pdf":
        import pymupdf as fitz
        assert fitz.open(path).page_count > 0
    elif ext == "docx":
        import docx
        docx.Document(path)
    elif ext == "pptx":
        from pptx import Presentation
        assert len(Presentation(path).slides) > 0
    elif ext == "xlsx":
        import openpyxl
        assert openpyxl.load_workbook(path).sheetnames
    elif ext in ("odt", "ods", "odp"):
        with zipfile.ZipFile(path) as z:
            assert z.read("mimetype").startswith(b"application/vnd.oasis.opendocument")
    elif ext in converters.IMAGES:
        from PIL import Image
        from pillow_heif import register_heif_opener
        register_heif_opener()
        with Image.open(path) as im:
            im.load()
            assert im.width > 0
    elif ext == "csv":
        import pandas as pd
        assert len(pd.read_csv(path).columns) > 0
    elif ext == "json":
        assert json.loads(path.read_text(encoding="utf-8"))
    elif ext in ("txt", "html", "md"):
        assert path.read_text(encoding="utf-8").strip()
    elif ext in converters.VIDEO + converters.AUDIO:
        run([find("ffmpeg"), "-v", "error", "-i", path, "-f", "null", "-"], 60)  # decodes the whole file
    else:
        pytest.fail(f"No check for .{ext}")


CASES = [(ext, out) for ext, outs in converters.MAP.items() for out in outs]


@pytest.mark.parametrize("ext,target", CASES, ids=[f"{e}->{t}" for e, t in CASES])
def test_conversion(ext, target, samples, tmp_path):
    missing = [TOOLS[t]["name"] for t in converters.MAP[ext][target].needs if not TOOLS[t]["found"]]
    if missing:
        pytest.skip(f"needs {', '.join(missing)}")
    src = samples / f"sample.{ext}"
    if not src.exists():
        pytest.skip(f"no sample.{ext} (a program needed to create it is missing)")
    outs = converters.run(target, [src], tmp_path, {"_timeout": 300})
    assert outs
    for o in outs:
        check(Path(o))


def test_ocr_reads_text(samples, tmp_path):
    if not TOOLS["tesseract"]["found"]:
        pytest.skip("needs Tesseract OCR")
    out = converters.run("ocr-txt", [samples / "sample.png"], tmp_path, {"_timeout": 120})[0]
    assert "12345" in out.read_text(encoding="utf-8")


def test_merge_pdfs(samples, tmp_path):
    out = converters.run("merge-pdf", [samples / "sample.pdf"] * 3, tmp_path, {})[0]
    import pymupdf as fitz
    assert fitz.open(out).page_count == 6


def test_images_to_one_pdf(samples, tmp_path):
    out = converters.run("images-pdf", [samples / "sample.png", samples / "sample.jpg", samples / "sample.heic"],
                         tmp_path, {})[0]
    import pymupdf as fitz
    assert fitz.open(out).page_count == 3


def test_split_with_page_range(samples, tmp_path):
    outs = converters.run("pdf-split", [samples / "sample.pdf"], tmp_path, {"pages": "2"})
    assert [o.name for o in outs] == ["sample_page2.pdf"]


def test_bad_page_range(samples, tmp_path):
    with pytest.raises(ConversionError, match="outside"):
        converters.run("txt", [samples / "sample.pdf"], tmp_path, {"pages": "5-9"})


def test_damaged_file_gives_clear_error(tmp_path):
    bad = tmp_path / "broken.png"
    bad.write_bytes(b"not an image")
    with pytest.raises(ConversionError, match="could not be opened"):
        converters.run("jpg", [bad], tmp_path, {})
