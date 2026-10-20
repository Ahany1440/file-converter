"""PDF work with PyMuPDF (text, html, images, merge, split) and pdf2docx / python-pptx."""
import html
from pathlib import Path

import pymupdf as fitz

from .tools import ConversionError


def parse_range(spec, page_count):
    """'1-3, 5' -> [0, 1, 2, 4]. Empty -> every page."""
    spec = (spec or "").strip()
    if not spec:
        return list(range(page_count))
    pages = []
    for part in spec.replace(" ", "").split(","):
        try:
            a, _, b = part.partition("-")
            first, last = int(a), int(b or a)
        except ValueError:
            raise ConversionError(f"Page range '{spec}' isn't valid. Use a form like 1-3, 5.")
        if not 1 <= first <= last <= page_count:
            raise ConversionError(f"Page range '{part}' is outside this document (it has {page_count} pages).")
        pages.extend(range(first - 1, last))
    return pages


def _open(src):
    try:
        doc = fitz.open(src)
    except Exception:
        raise ConversionError("This PDF could not be opened (it may be damaged).")
    if doc.needs_pass:
        raise ConversionError("This PDF is password-protected. Remove the password first.")
    return doc


def to_text(src, out_dir, fmt, opts):
    doc = _open(src)
    out = Path(out_dir) / f"{Path(src).stem}.txt"
    out.write_text("\n\f".join(doc[i].get_text() for i in parse_range(opts.get("pages"), len(doc))), encoding="utf-8")
    return [out]


def to_html(src, out_dir, fmt, opts):
    doc = _open(src)
    body = "\n".join(doc[i].get_text("html") for i in parse_range(opts.get("pages"), len(doc)))
    out = Path(out_dir) / f"{Path(src).stem}.html"
    out.write_text(f'<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(Path(src).stem)}</title>'
                   f"</head><body>{body}</body></html>", encoding="utf-8")
    return [out]


def to_images(src, out_dir, fmt, opts):
    doc = _open(src)
    dpi = int(opts.get("dpi") or 150)
    stem, outs = Path(src).stem, []
    for i in parse_range(opts.get("pages"), len(doc)):
        pix = doc[i].get_pixmap(dpi=dpi)
        out = Path(out_dir) / f"{stem}_page{i + 1}.{fmt}"
        if fmt == "jpg":
            pix.save(out, jpg_quality=int(opts.get("quality") or 90))
        else:
            pix.save(out)
        outs.append(out)
    return outs


def to_docx(src, out_dir, fmt, opts):
    from pdf2docx import Converter
    doc = _open(src)
    pages = parse_range(opts.get("pages"), len(doc))
    doc.close()
    out = Path(out_dir) / f"{Path(src).stem}.docx"
    cv = Converter(str(src))
    try:
        cv.convert(str(out), pages=pages)
    finally:
        cv.close()
    return [out]


def to_pptx(src, out_dir, fmt, opts):
    """One slide per page, each slide a full-size picture of the page."""
    from pptx import Presentation
    from pptx.util import Emu
    images = to_images(src, out_dir, "png", opts)
    first = fitz.open(images[0])[0].rect
    prs = Presentation()
    prs.slide_width = Emu(12192000)  # 13.333 in, the 16:9 default width
    prs.slide_height = Emu(int(12192000 * first.height / first.width))
    for img in images:
        prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(str(img), 0, 0, prs.slide_width, prs.slide_height)
        img.unlink()
    out = Path(out_dir) / f"{Path(src).stem}.pptx"
    prs.save(out)
    return [out]


def split(src, out_dir, fmt, opts):
    doc = _open(src)
    stem, outs = Path(src).stem, []
    for i in parse_range(opts.get("pages"), len(doc)):
        page = fitz.open()
        page.insert_pdf(doc, from_page=i, to_page=i)
        out = Path(out_dir) / f"{stem}_page{i + 1}.pdf"
        page.save(out)
        outs.append(out)
    return outs


def merge(srcs, out_dir, opts):
    merged = fitz.open()
    for src in srcs:
        merged.insert_pdf(_open(src))
    out = Path(out_dir) / "merged.pdf"
    merged.save(out, garbage=3, deflate=True)
    return [out]
